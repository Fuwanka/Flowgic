"""Rebuild trusted templates with python-docx; run with bundled artifact Python."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent


def set_cell_borders(cell):
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        item = OxmlElement(f'w:{edge}')
        item.set(qn('w:val'), 'single')
        item.set(qn('w:sz'), '4')
        item.set(qn('w:color'), 'D9D9D9')
        borders.append(item)
    properties.append(borders)
    margins = OxmlElement('w:tcMar')
    for side in ('top', 'left', 'bottom', 'right'):
        item = OxmlElement(f'w:{side}')
        item.set(qn('w:w'), '85' if side in ('left', 'right') else '65')
        item.set(qn('w:type'), 'dxa')
        margins.append(item)
    properties.append(margins)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.4), Cm(1.3)
    section.left_margin, section.right_margin = Cm(1.7), Cm(1.7)
    section.header_distance, section.footer_distance = Cm(0.6), Cm(0.6)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
        # The bundled Word starter carries a blue bottom border on Title.
        # These forms intentionally use plain black typographic headings.
        for border in list(style.element.get_or_add_pPr().findall(qn('w:pBdr'))):
            border.getparent().remove(border)
        for spacing in list(style.element.get_or_add_rPr().findall(qn('w:spacing'))):
            spacing.getparent().remove(spacing)
        style.paragraph_format.space_after = Pt(5)
    normal = doc.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.05
    for name in ('Title', 'Subtitle'):
        doc.styles[name].paragraph_format.keep_with_next = True
    doc.styles['Title'].font.size = Pt(17)
    doc.styles['Title'].font.bold = True
    doc.styles['Subtitle'].font.size = Pt(10)
    doc.styles['Subtitle'].font.italic = False
    doc.styles['Heading 1'].font.size = Pt(11)
    doc.styles['Heading 1'].font.bold = True
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
    doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)
    doc.styles['Heading 1'].paragraph_format.keep_with_next = True
    doc.add_paragraph(title, style='Title')
    doc.add_paragraph(subtitle, style='Subtitle')
    header = section.header.paragraphs[0]
    header.text = 'Flowgic   Документы перевозки'
    header.style = doc.styles['Normal']
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.text = 'Проект документа   Заказ {{ order_number }}'
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(8)
    doc.core_properties.author = 'Flowgic'
    doc.core_properties.last_modified_by = 'Flowgic'
    doc.core_properties.title = title
    doc.core_properties.subject = 'Проект документа по данным заказа'
    return doc


def records(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width, table.columns[1].width = Cm(4.1), Cm(13.5)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = Cm(4.1), Cm(13.5)
        cells[0].text, cells[1].text = label, value
        for i, cell in enumerate(cells):
            set_cell_borders(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.keep_with_next = False
                if i == 0:
                    for run in paragraph.runs:
                        run.font.bold = True
            if i == 0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F2F2F2')
                cell._tc.get_or_add_tcPr().append(shading)
        cant_split = OxmlElement('w:cantSplit')
        table.rows[-1]._tr.get_or_add_trPr().append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def application():
    doc = document('Договор-заявка на перевозку', 'Договор-заявка № {{ order_number }} от {{ issued_date }}')
    doc.add_paragraph('ЧЕРНОВИК. Параметры перевозки взяты из заказа. Перед подписанием стороны проверяют реквизиты и дополняют условия, отсутствующие в карточке.')
    doc.add_heading('Стороны', level=1)
    records(doc, [
        ('Перевозчик', '{{ company_name }}\nИНН {{ company_inn }}'),
        ('Адрес перевозчика', '{{ company_address }}'),
        ('Заказчик', '{{ client_name }}'),
        ('Контакты заказчика', 'Телефон {{ client_phone }}   Email {{ client_email }}'),
    ])
    doc.add_heading('Перевозка', level=1)
    records(doc, [
        ('Погрузка', '{{ origin }}\nПлан {{ pickup_datetime }}'),
        ('Выгрузка', '{{ destination }}\nПлан {{ delivery_datetime }}'),
        ('Груз и масса', '{{ cargo_type }}   Масса {{ cargo_mass_kg }} кг'),
        ('Расстояние', '{{ distance_km }} км'),
    ])
    doc.add_heading('Ресурсы и стоимость', level=1)
    records(doc, [
        ('Транспорт', '{{ vehicle_number }}   {{ vehicle_model }}\nТип {{ vehicle_type }}   Грузоподъёмность {{ vehicle_capacity }} кг'),
        ('Водитель', '{{ driver_name }}   Телефон {{ driver_phone }}'),
        ('Стоимость перевозки', '{{ agreed_price }} руб.   НДС и порядок оплаты не указаны'),
    ])
    doc.add_paragraph('Время указано в часовом поясе {{ timezone_name }}. Статус заказа {{ status }}. Особые условия, ответственность сторон и банковские реквизиты заполняются и согласуются отдельно.')
    doc.add_heading('Согласование', level=1)
    doc.add_paragraph('От перевозчика ____________________ / ____________________')
    doc.add_paragraph('От заказчика ______________________ / ____________________')
    doc.add_paragraph('Дата подписания __________________')
    doc.save(HERE / 'application.docx')


def ttn():
    doc = document('Товарно-транспортная накладная', 'ТТН по заказу № {{ order_number }} от {{ issued_date }}')
    doc.add_paragraph('ЧЕРНОВИК. Рабочий проект по данным заказа. Реквизиты грузоотправителя, грузополучателя, упаковки и фактической приёмки дополняются сторонами.')
    doc.add_heading('Участники и маршрут', level=1)
    records(doc, [
        ('Перевозчик', '{{ company_name }}   ИНН {{ company_inn }}'),
        ('Адрес перевозчика', '{{ company_address }}'),
        ('Заказчик перевозки', '{{ client_name }}\nТелефон {{ client_phone }}   Email {{ client_email }}'),
        ('Грузоотправитель', 'Не указан   Реквизиты ______________________________'),
        ('Грузополучатель', 'Не указан   Реквизиты ______________________________'),
        ('Пункт погрузки', '{{ origin }}\nПлан {{ pickup_datetime }}'),
        ('Пункт выгрузки', '{{ destination }}\nПлан {{ delivery_datetime }}'),
    ])
    doc.add_heading('Груз', level=1)
    records(doc, [
        ('Наименование', '{{ cargo_type }}'),
        ('Масса по заказу', '{{ cargo_mass_kg }} кг'),
        ('Упаковка и места', 'Не указаны   Вид __________   Количество __________'),
    ])
    doc.add_heading('Транспортные данные', level=1)
    records(doc, [
        ('Автомобиль', '{{ vehicle_number }}   {{ vehicle_model }}\nТип {{ vehicle_type }}   Грузоподъёмность {{ vehicle_capacity }} кг'),
        ('Водитель', '{{ driver_name }}   Телефон {{ driver_phone }}'),
        ('Перевозка', '{{ distance_km }} км   Стоимость {{ agreed_price }} руб.'),
    ])
    doc.add_paragraph('Время указано в часовом поясе {{ timezone_name }}. Фактические даты, масса, состояние груза и расхождения отмечаются при передаче.')
    doc.add_heading('Передача груза', level=1)
    doc.add_paragraph('Отпустил __________________   Принял водитель __________________')
    doc.add_paragraph('Сдал водитель ______________   Получил __________________________')
    doc.add_paragraph('Дата и отметки _______________________________________________')
    doc.save(HERE / 'ttn.docx')


if __name__ == '__main__':
    application()
    ttn()
    print('Created application.docx and ttn.docx')
